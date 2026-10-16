# Copyright (C) [2025] [阿炜]
# 本程序是自由软件：你可以根据 GNU GPL v3 协议重新分发或修改它。
"""Excel 导出算法对照测试。

用一份已知场景的临时库，断言 collect_export_rows 的输出矩阵与预期完全一致，
并验证合并单元格区间符合旧实现（含那个刻意保留的运算符优先级行为）。

运行：python -m unittest discover -s tests -v
"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.core.db import ReagentDatabase  # noqa: E402
from app.core import excel  # noqa: E402

START = "2025-02-01"
END = "2025-03-31"


def build_fixture_db():
    """构造覆盖四种情形的库：

    A/A1 跨越窗口（窗口前有记录 + 窗口内有出入库）
    B/B1 仅窗口前有记录，窗口内无变动      -> 应补一条占位行
    C/C1 仅窗口内有记录                    -> 历史结余为 0
    D/D1 窗口前已清零且窗口内无变动        -> 应被过滤掉
    """
    path = os.path.join(tempfile.mkdtemp(), "fixture.db")
    db = ReagentDatabase(path)

    def setup(name, batch):
        db.add_new_reagent(name, batch, "2024-12-01", "2026-12-01", 0)
        return db.get_reagent_details(name, batch)[0]

    # A/A1
    rid = setup("试剂A", "A1")
    db.add_inbound(rid, "试剂A", "A1", "2025-01-05 09:00:00:000", 100, "张三")
    db.add_inbound(rid, "试剂A", "A1", "2025-02-10 09:00:00:000", 50, "张三")
    db.add_outbound(rid, "试剂A", "A1", "2025-03-01 09:00:00:000", 30, "李四")

    # B/B1
    rid = setup("试剂B", "B1")
    db.add_inbound(rid, "试剂B", "B1", "2025-01-01 09:00:00:000", 20, "张三")

    # C/C1
    rid = setup("试剂C", "C1")
    db.add_inbound(rid, "试剂C", "C1", "2025-02-20 09:00:00:000", 10, "王五")

    # D/D1 —— 窗口前入库又全部出库，结余为 0
    rid = setup("试剂D", "D1")
    db.add_inbound(rid, "试剂D", "D1", "2025-01-10 09:00:00:000", 5, "张三")
    db.add_outbound(rid, "试剂D", "D1", "2025-01-20 09:00:00:000", 5, "张三")

    return db


class TestExportRows(unittest.TestCase):
    def setUp(self):
        self.db = build_fixture_db()

    def test_header_is_exact(self):
        rows = excel.collect_export_rows(self.db, START, END)
        self.assertEqual(rows[0], [
            "试剂名称", "试剂批号", "历史结余", "入库日期", "入库数量",
            "出库日期", "出库数量", "操作后库存", "周期结余", "库存校验",
        ])

    def test_body_matches_expected(self):
        rows = excel.collect_export_rows(self.db, START, END)
        body = sorted(rows[1:], key=lambda r: (r[0], r[1], str(r[3]), str(r[5])))

        expected = sorted([
            # A/A1：跨越窗口
            ["试剂A", "A1", 100, "2025-02-10 09:00:00:000", 50, "", "", 150, 120, ""],
            ["试剂A", "A1", 100, "", "", "2025-03-01 09:00:00:000", 30, 120, 120, ""],
            # B/B1：窗口内无变动 -> 占位行，历史/周期结余均为 20
            ["试剂B", "B1", 20, "", "", "", "", "", 20, ""],
            # C/C1：窗口内新建，历史结余 0
            ["试剂C", "C1", 0, "2025-02-20 09:00:00:000", 10, "", "", 10, 10, ""],
        ], key=lambda r: (r[0], r[1], str(r[3]), str(r[5])))

        self.assertEqual(body, expected)

    def test_zero_balance_batch_is_filtered_out(self):
        """D/D1 三项皆为零，不应出现在导出结果里。"""
        rows = excel.collect_export_rows(self.db, START, END)
        names = {r[0] for r in rows[1:]}
        self.assertNotIn("试剂D", names)
        self.assertEqual(len(rows), 1 + 4, "应为表头 + 4 行数据")


class TestWorkbookOutput(unittest.TestCase):
    def setUp(self):
        self.db = build_fixture_db()
        self.out = os.path.join(tempfile.mkdtemp(), "out.xlsx")

    def test_written_file_has_expected_values_and_merges(self):
        from openpyxl import load_workbook

        rows = excel.collect_export_rows(self.db, START, END)
        excel.write_workbook(rows, self.out)

        wb = load_workbook(self.out)
        ws = wb[excel.SHEET_TITLE]
        self.assertEqual(ws.max_row, len(rows))
        self.assertEqual(ws.max_column, 10)

        # 表头加粗
        self.assertTrue(ws.cell(row=1, column=1).font.bold)

        # 数值抽查：A/A1 的历史结余与周期结余被纵向合并，取值仍来自首行
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertTrue(any(r.startswith("B2:") for r in merged), f"A/A1 批号列应合并: {merged}")

    def test_default_filename_uses_two_digit_year(self):
        self.assertEqual(
            excel.default_filename("2025-01-01", "2025-06-30"),
            "出入库记录_25-01-01_至_25-06-30.xlsx",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)

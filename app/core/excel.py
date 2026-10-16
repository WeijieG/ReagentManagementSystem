# Copyright (C) [2025] [阿炜]
# 本程序是自由软件：你可以根据 GNU GPL v3 协议重新分发或修改它。
"""出入库记录导出为 XLSX。

算法照搬自 ReagentMGS.py 的 MoreDialog.export_records_to_xlsx（原文件 2105-2396 行），
未做行为改动。仅拆成两段：

  collect_export_rows()  —— 纯取数，返回二维表，可脱离 Qt 测试
  write_workbook()       —— 只负责写盘与样式

这样后台线程可以只承担写盘，避免 SQLite 连接跨线程。
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "试剂名称", "试剂批号", "历史结余", "入库日期", "入库数量",
    "出库日期", "出库数量", "操作后库存", "周期结余", "库存校验",
]

COLUMN_WIDTHS = [20, 20, 10, 15, 10, 15, 10, 10, 10, 10]

SHEET_TITLE = "出入库记录"

# 周期末日期的后缀。与旧实现一致：开始日期不加后缀（含当天字面前缀即可），
# 结束日期加后缀以覆盖带时间戳的记录。
END_OF_DAY_SUFFIX = " 23:59:59:999"


def default_filename(start_date, end_date):
    """生成默认文件名，保持旧实现的两位数年份写法。"""
    short = lambda d: d[2:] if len(d) >= 8 else d  # noqa: E731  "2025-01-01" -> "25-01-01"
    return f"出入库记录_{short(start_date)}_至_{short(end_date)}.xlsx"


def _distinct_reagents(db, end_date_q):
    """在结束日期前有记录的所有 (名称, 批号)。"""
    db.cursor.execute('''
        SELECT DISTINCT reagent_name, reagent_batch
        FROM (
            SELECT reagent_name, reagent_batch
            FROM inbound
            WHERE date <= ?
            UNION
            SELECT reagent_name, reagent_batch
            FROM outbound
            WHERE date <= ?
        )
    ''', (end_date_q, end_date_q))
    return db.cursor.fetchall()


def _stock_at(db, name, batch, boundary, inclusive):
    """取某 (名称, 批号) 在边界处的最后一条记录的操作后库存。

    inclusive=True 时用 date <= boundary（周期末），否则 date < boundary（周期开始前）。
    """
    operator = "<=" if inclusive else "<"
    db.cursor.execute(f'''
        SELECT remaining_quantity, date, 'in' as type
        FROM inbound
        WHERE reagent_name = ? AND reagent_batch = ? AND date {operator} ?
        UNION
        SELECT remaining_quantity, date, 'out' as type
        FROM outbound
        WHERE reagent_name = ? AND reagent_batch = ? AND date {operator} ?
        ORDER BY date DESC
        LIMIT 1
    ''', (name, batch, boundary, name, batch, boundary))
    row = db.cursor.fetchone()
    return row[0] if row else 0


def collect_export_rows(db, start_date, end_date):
    """按旧实现的口径收集导出数据。

    start_date / end_date 均为 'yyyy-MM-dd'。返回二维列表，首行为表头。
    """
    end_date_q = end_date + END_OF_DAY_SUFFIX

    distinct_reagents = _distinct_reagents(db, end_date_q)

    # 1. 选定周期内的记录，合并并打上类型标记
    all_records = []
    for rec in db.get_inbound_records("", "", start_date, end_date_q):
        # (id, reagent_name, reagent_batch, date, quantity, operator, remaining_quantity)
        all_records.append(('in', rec[1], rec[2], rec[3], rec[4], rec[5], rec[6]))
    for rec in db.get_outbound_records("", "", start_date, end_date_q):
        all_records.append(('out', rec[1], rec[2], rec[3], rec[4], rec[5], rec[6]))

    all_records.sort(key=lambda x: (x[1], x[2], x[3]))

    # 2. 每个批号准备一个槽位
    organized_data = {}
    for name, batch in distinct_reagents:
        organized_data[(name, batch)] = {
            'initial_stock': 0,
            'records': [],
            'ending_stock': 0,
        }

    # 3. 历史结余：周期开始前最后一条记录
    for (name, batch), data in organized_data.items():
        data['initial_stock'] = _stock_at(db, name, batch, start_date, inclusive=False)

    # 4. 周期结余：周期结束前最后一条记录
    for (name, batch), data in organized_data.items():
        data['ending_stock'] = _stock_at(db, name, batch, end_date_q, inclusive=True)

    # 5. 挂上周期内记录
    for record in all_records:
        record_type, name, batch, date, quantity, operator, remaining = record
        key = (name, batch)
        if key in organized_data:
            organized_data[key]['records'].append({
                'type': record_type,
                'date': date,
                'quantity': quantity,
                'operator': operator,
                'remaining': remaining,
            })

    # 6. 过滤：历史结余为 0、周期内无记录、且周期结余为 0 的批号不导出
    filtered_data = {}
    for key, data in organized_data.items():
        if data['initial_stock'] == 0 and len(data['records']) == 0 and data['ending_stock'] == 0:
            continue
        filtered_data[key] = data

    # 7. 铺平成二维表
    excel_data = [list(HEADERS)]

    for (name, batch), data in filtered_data.items():
        # 周期内无记录的批号补一个空行，让历史结余/周期结余仍能打印出来
        if len(data['records']) == 0:
            data['records'].append({
                'type': 'in', 'date': '', 'quantity': '', 'operator': '', 'remaining': '',
            })

        for record in data['records']:
            if record['type'] == 'in':
                excel_data.append([
                    name, batch, data['initial_stock'],
                    record['date'], record['quantity'],
                    "", "",
                    record['remaining'],
                    data['ending_stock'], "",
                ])
            else:
                excel_data.append([
                    name, batch, data['initial_stock'],
                    "", "",
                    record['date'], record['quantity'],
                    record['remaining'],
                    data['ending_stock'], "",
                ])

    return excel_data


def write_workbook(excel_data, file_path):
    """把 collect_export_rows 的结果写成 xlsx 并应用样式与合并。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    for row in excel_data:
        ws.append(row)

    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    center_aligned = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows(min_row=1, max_row=len(excel_data)):
        for cell in row:
            cell.alignment = center_aligned

    _merge_cells(ws, excel_data)

    wb.save(file_path)
    return file_path


def _merge_cells(ws, excel_data):
    """纵向合并同一 (名称, 批号) 分组内的 B/C/I/J 列，以及同名分组的 A 列。

    注意：下面这一行的条件写法是**照搬旧实现**（原文件 2356 行）：
        if (batch != current_batch) or (name != current_name) and row_idx > batch_start_row:

    由于 Python 中 `and` 优先级高于 `or`，它实际等价于：
        (batch != current_batch) or ((name != current_name) and row_idx > batch_start_row)

    也就是说，当名称变化但批号相同时，若 row_idx == batch_start_row，前一组
    不会在这里收口。这几乎肯定不是原意（原意应当是两边都加括号）。

    本次重构**刻意保留**该行为，以保证导出结果与旧版逐单元格、逐合并区间一致。
    是否修正属于独立决策，需要你确认后再动。
    """
    current_name = None
    name_start_row = 2
    current_batch = None
    batch_start_row = 2

    for row_idx in range(2, len(excel_data) + 1):
        name = ws.cell(row=row_idx, column=1).value
        batch = ws.cell(row=row_idx, column=2).value

        if not current_name:
            current_name = name
        if not current_batch:
            current_batch = batch

        # 合并批号分组（保留旧实现的运算符优先级）
        if name and batch:
            if (batch != current_batch) or (name != current_name) and row_idx > batch_start_row:
                ws.merge_cells(f'B{batch_start_row}:B{row_idx - 1}')
                ws.merge_cells(f'C{batch_start_row}:C{row_idx - 1}')
                ws.merge_cells(f'I{batch_start_row}:I{row_idx - 1}')
                ws.merge_cells(f'J{batch_start_row}:J{row_idx - 1}')
                current_batch = batch
                batch_start_row = row_idx

        # 合并名称分组
        if name:
            if name != current_name and row_idx > name_start_row:
                ws.merge_cells(f'A{name_start_row}:A{row_idx - 1}')
                current_name = name
                name_start_row = row_idx

        # 最后一行收口
        if row_idx == len(excel_data):
            if name == current_name:
                ws.merge_cells(f'A{name_start_row}:A{row_idx}')
            if batch == current_batch:
                ws.merge_cells(f'B{batch_start_row}:B{row_idx}')
                ws.merge_cells(f'C{batch_start_row}:C{row_idx}')
                ws.merge_cells(f'I{batch_start_row}:I{row_idx}')
                ws.merge_cells(f'J{batch_start_row}:J{row_idx}')
